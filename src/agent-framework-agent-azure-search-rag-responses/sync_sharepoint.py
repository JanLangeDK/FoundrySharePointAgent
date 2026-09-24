import asyncio
import hashlib
import io
import json
import logging
import os
import re
from collections.abc import AsyncIterator, Iterable
from html.parser import HTMLParser
from pathlib import Path
from typing import Any
from urllib.parse import quote, unquote, urlparse

import httpx
from azure.identity.aio import DefaultAzureCredential
from azure.search.documents.aio import SearchClient
from docx import Document
from dotenv import load_dotenv
from openpyxl import load_workbook
from pypdf import PdfReader
from pptx import Presentation

from sharepoint_source import SharePointSource, parse_sharepoint_folder_url

GRAPH_ROOT = "https://graph.microsoft.com/v1.0"
GRAPH_SCOPE = "https://graph.microsoft.com/.default"
SUPPORTED_EXTENSIONS = {".csv", ".docx", ".html", ".json", ".md", ".pdf", ".pptx", ".txt", ".xlsx", ".xml"}

logger = logging.getLogger("sharepoint_sync")


class GraphRequestError(RuntimeError):
    pass


class _TextExtractor(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []

    def handle_data(self, data: str) -> None:
        if data.strip():
            self.parts.append(data.strip())


class GraphClient:
    def __init__(self, credential: DefaultAzureCredential) -> None:
        self._credential = credential
        self._client = httpx.AsyncClient(follow_redirects=True, timeout=60)

    async def close(self) -> None:
        await self._client.aclose()

    async def _headers(self) -> dict[str, str]:
        token = await self._credential.get_token(GRAPH_SCOPE)
        return {"Authorization": f"Bearer {token.token}"}

    async def get_json(self, path_or_url: str) -> dict[str, Any]:
        url = path_or_url if path_or_url.startswith("https://") else f"{GRAPH_ROOT}{path_or_url}"
        response = await self._client.get(url, headers=await self._headers())
        if response.is_error:
            raise GraphRequestError(f"Microsoft Graph {response.status_code} for {url}: {response.text}")
        return response.json()

    async def iter_values(self, path_or_url: str) -> AsyncIterator[dict[str, Any]]:
        next_url: str | None = path_or_url
        while next_url:
            page = await self.get_json(next_url)
            for value in page.get("value", []):
                yield value
            next_url = page.get("@odata.nextLink")

    async def download(self, drive_id: str, item_id: str) -> bytes:
        url = f"{GRAPH_ROOT}/drives/{quote(drive_id, safe='')}/items/{quote(item_id, safe='')}/content"
        response = await self._client.get(url, headers=await self._headers())
        if response.is_error:
            raise GraphRequestError(f"Microsoft Graph {response.status_code} downloading item {item_id}")
        return response.content


def extract_allowed_principals(permissions: Iterable[dict[str, Any]], tenant_id: str) -> set[str]:
    principals: set[str] = set()
    for permission in permissions:
        if not set(permission.get("roles", [])) & {"read", "write", "owner"}:
            continue

        identity_sets = [permission.get("grantedToV2", {})]
        identity_sets.extend(permission.get("grantedToIdentitiesV2", []))
        for identity_set in identity_sets:
            for identity_type in ("user", "group"):
                principal_id = identity_set.get(identity_type, {}).get("id")
                if principal_id:
                    principals.add(str(principal_id))
            site_group_id = identity_set.get("siteGroup", {}).get("id")
            if site_group_id:
                principals.add(f"siteGroup:{site_group_id}")

        link_scope = permission.get("link", {}).get("scope")
        if link_scope == "organization":
            principals.add(f"tenant:{tenant_id}")
    return principals


def extract_text(name: str, content: bytes) -> str:
    extension = Path(name).suffix.lower()
    if extension in {".csv", ".json", ".md", ".txt", ".xml"}:
        return content.decode("utf-8-sig", errors="replace")
    if extension == ".html":
        parser = _TextExtractor()
        parser.feed(content.decode("utf-8-sig", errors="replace"))
        return "\n".join(parser.parts)
    if extension == ".pdf":
        return "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(content)).pages)
    if extension == ".docx":
        document = Document(io.BytesIO(content))
        parts = [paragraph.text for paragraph in document.paragraphs]
        parts.extend("\t".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows)
        return "\n".join(parts)
    if extension == ".pptx":
        presentation = Presentation(io.BytesIO(content))
        return "\n".join(
            shape.text
            for slide in presentation.slides
            for shape in slide.shapes
            if hasattr(shape, "text") and shape.text.strip()
        )
    if extension == ".xlsx":
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            return "\n".join(
                "\t".join(str(value) for value in row if value is not None)
                for worksheet in workbook.worksheets
                for row in worksheet.iter_rows(values_only=True)
                if any(value is not None for value in row)
            )
        finally:
            workbook.close()
    raise ValueError(f"Unsupported file extension: {extension}")


def chunk_text(text: str, max_characters: int = 6000, overlap_characters: int = 500) -> list[str]:
    normalized = re.sub(r"[ \t]+", " ", text)
    normalized = re.sub(r"\n{3,}", "\n\n", normalized).strip()
    if not normalized:
        return []

    chunks: list[str] = []
    start = 0
    while start < len(normalized):
        end = min(start + max_characters, len(normalized))
        if end < len(normalized):
            boundary = max(normalized.rfind("\n", start, end), normalized.rfind(" ", start, end))
            if boundary > start + max_characters // 2:
                end = boundary
        chunks.append(normalized[start:end].strip())
        if end == len(normalized):
            break
        start = max(end - overlap_characters, start + 1)
    return chunks


async def resolve_site_and_drive(graph: GraphClient, source: SharePointSource) -> tuple[str, str]:
    site = await graph.get_json("/sites/root?$select=id,webUrl")
    if urlparse(site["webUrl"]).hostname != source.hostname:
        raise GraphRequestError(f"Authenticated tenant root site does not match {source.hostname}.")

    site_id = site["id"]
    visible_drive_count = 0
    async for drive in graph.iter_values(f"/sites/{quote(site_id, safe='')}/drives?$select=id,name,webUrl"):
        visible_drive_count += 1
        drive_path = unquote(urlparse(drive["webUrl"]).path).rstrip("/")
        if drive_path == source.library_path.rstrip("/"):
            return site_id, drive["id"]
    if visible_drive_count == 0:
        raise GraphRequestError(
            "Microsoft Graph returned no visible document libraries. Grant the ingestion identity "
            "Sites.Selected access to this site (or a broader Files/Sites read permission)."
        )
    raise GraphRequestError(f"No document library matched {source.library_path!r}.")


async def iter_drive_items(graph: GraphClient, drive_id: str, folder_path: str) -> AsyncIterator[dict[str, Any]]:
    fields = "id,name,webUrl,size,eTag,file,folder"
    if folder_path:
        encoded_path = quote(folder_path, safe="/")
        first_url = f"/drives/{quote(drive_id, safe='')}/root:/{encoded_path}:/children?$select={fields}"
    else:
        first_url = f"/drives/{quote(drive_id, safe='')}/root/children?$select={fields}"

    pending_urls = [first_url]
    while pending_urls:
        current_url = pending_urls.pop()
        async for item in graph.iter_values(current_url):
            if item.get("folder"):
                pending_urls.append(
                    f"/drives/{quote(drive_id, safe='')}/items/{quote(item['id'], safe='')}/children?$select={fields}"
                )
            elif item.get("file"):
                yield item


async def get_item_principals(
    graph: GraphClient, drive_id: str, item_id: str, tenant_id: str
) -> set[str]:
    permissions_url = (
        f"/drives/{quote(drive_id, safe='')}/items/{quote(item_id, safe='')}/permissions"
        "?$select=id,roles,grantedToV2,grantedToIdentitiesV2,link,inheritedFrom"
    )
    permissions = [permission async for permission in graph.iter_values(permissions_url)]
    return extract_allowed_principals(permissions, tenant_id)


async def existing_source_document_ids(search: SearchClient, source_root: str) -> set[str]:
    escaped_source = source_root.replace("'", "''")
    results = await search.search(
        search_text="*",
        filter=f"sourceRoot eq '{escaped_source}'",
        select=["id"],
    )
    return {document["id"] async for document in results}


async def sync() -> None:
    load_dotenv()
    source_url = os.environ["SHAREPOINT_FOLDER_URL"]
    source = parse_sharepoint_folder_url(source_url)
    search_endpoint = os.environ["AZURE_SEARCH_ENDPOINT"]
    search_index_name = os.environ.get("AZURE_SEARCH_INDEX_NAME", "sharepoint-documents")
    tenant_id = os.environ["AZURE_TENANT_ID"]
    max_file_bytes = int(os.environ.get("SHAREPOINT_MAX_FILE_BYTES", str(50 * 1024 * 1024)))

    indexed_documents: list[dict[str, Any]] = []
    skipped_files = 0

    async with DefaultAzureCredential() as credential:
        graph = GraphClient(credential)
        try:
            _, drive_id = await resolve_site_and_drive(graph, source)
            async with SearchClient(search_endpoint, search_index_name, credential) as search:
                previous_ids = await existing_source_document_ids(search, source_url)
                async for item in iter_drive_items(graph, drive_id, source.folder_path):
                    extension = Path(item["name"]).suffix.lower()
                    if extension not in SUPPORTED_EXTENSIONS or item.get("size", 0) > max_file_bytes:
                        skipped_files += 1
                        continue

                    principals = await get_item_principals(graph, drive_id, item["id"], tenant_id)
                    if not principals:
                        logger.warning("Skipping %s because no supported effective principals were found.", item["name"])
                        skipped_files += 1
                        continue

                    content = await graph.download(drive_id, item["id"])
                    for chunk_index, chunk in enumerate(chunk_text(extract_text(item["name"], content))):
                        document_id = hashlib.sha256(f"{drive_id}:{item['id']}:{chunk_index}".encode()).hexdigest()
                        indexed_documents.append(
                            {
                                "id": document_id,
                                "content": chunk,
                                "sourceName": item["name"],
                                "sourceLink": item["webUrl"],
                                "sourceItemId": item["id"],
                                "sourceETag": item.get("eTag", ""),
                                "sourceRoot": source_url,
                                "chunkIndex": chunk_index,
                                "allowedPrincipals": sorted(principals),
                            }
                        )

                current_ids = {document["id"] for document in indexed_documents}
                for offset in range(0, len(indexed_documents), 500):
                    results = await search.merge_or_upload_documents(indexed_documents[offset : offset + 500])
                    failures = [(result.key, result.error_message) for result in results if not result.succeeded]
                    if failures:
                        raise RuntimeError(f"Azure AI Search upload failures: {failures}")

                stale_ids = sorted(previous_ids - current_ids)
                for offset in range(0, len(stale_ids), 500):
                    batch = [{"id": document_id} for document_id in stale_ids[offset : offset + 500]]
                    if batch:
                        await search.delete_documents(batch)
        finally:
            await graph.close()

    print(
        json.dumps(
            {
                "source": source_url,
                "indexedChunks": len(indexed_documents),
                "skippedFiles": skipped_files,
            },
            indent=2,
        )
    )


if __name__ == "__main__":
    logging.basicConfig(level=os.getenv("LOG_LEVEL", "INFO").upper())
    asyncio.run(sync())